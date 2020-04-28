# Standard library imports
import logging
from datetime import datetime

# Related third party imports
import xlsxwriter
import docx
from cerberus import Validator
from docx import Document
from flask_restful import current_app, Api
from flask import Blueprint, request, send_file
from xlsxwriter import Workbook
from xlsxwriter.worksheet import Worksheet

# Local application/library specific imports
from .base import BaseResource
from ..models.user import User as UserModel
from ..utils import to_readable, pos_to_char, find_longest_word
from ..libs.libreoffice import convert_to
from ..utils.cerberus_schema import user_model_schema, search_model_schema

blueprint = Blueprint('users', __name__, url_prefix='/users')
api = Api(blueprint)

logger = logging.getLogger(__name__)


class UserResource(BaseResource):
    db_model = UserModel
    column_display_order = ['name', 'last_name', 'age', 'birth_date', 'role', 'created_at', 'updated_at',
                            'deleted_at']

    def format_column_names(self, rows: list, original_column_names: set) -> None:
        formatted_column_names = [
            column.title().replace('_', ' ')
            for column in original_column_names
            if column
        ]

        rows.append(formatted_column_names)

    def get_users(self, column_names: list, page_number: int, items_per_page: int) -> UserModel:
        select_fields = [
            UserModel._meta.fields[column_name]
            for column_name in column_names
        ]

        users_query = (UserModel.select(*select_fields)
                       .paginate(page_number, items_per_page)
                       .dicts())

        return users_query

    def format_user_data(self, users_query: list, rows: list) -> None:
        users_list = []

        for user in users_query:
            user_dict = {
                'role': user.role.name,
            }

            user_dict.update({
                k: to_readable(v)
                for (k, v) in user.serialize(['id', 'role']).items()
            })

            user_dict = dict(sorted(user_dict.items(), key=lambda x: self.column_display_order.index(x[0])))

            users_list.append(user_dict)

        for user_dict in users_list:
            user_values = list(user_dict.values())
            rows.append(user_values)


@api.resource('')
class NewUserResource(UserResource):
    def post(self) -> tuple:
        data = request.get_json()

        v = Validator(schema=user_model_schema())
        v.allow_unknown = False

        if not v.validate(data):
            return {
                       'message': 'validation error',
                       'fields': v.errors,
                   }, 422

        user = UserModel.create(**data)
        user_dict = user.serialize()

        return {
                   'data': user_dict,
               }, 201


@api.resource('/<int:user_id>')
class UserResource(UserResource):
    def get(self, user_id: int) -> tuple:
        response = {
            'error': 'User doesn\'t exist',
        }
        status_code = 404

        user = UserModel.get_or_none(UserModel.id == user_id)

        if isinstance(user, UserModel):
            user_dict = user.serialize()

            response = {
                'data': user_dict,
            }
            status_code = 200

        return response, status_code

    def put(self, user_id: int) -> tuple:
        data = request.get_json()

        v = Validator(schema=user_model_schema())
        v.allow_unknown = False

        if not v.validate(data):
            return {
                       'message': 'validation error',
                       'fields': v.errors
                   }, 422

        user = (UserModel.get_or_none(UserModel.id == user_id,
                                      UserModel.deleted_at.is_null()))
        if user:
            data['id'] = user_id
            UserModel(**data).save()

            user = (UserModel.get_or_none(UserModel.id == user_id,
                                          UserModel.deleted_at.is_null()))
            user_dict = user.serialize()

            response_data = {
                'data': user_dict,
            }
            response_code = 200
        else:
            response_data = {
                'error': 'User doesn\'t exist',
            }
            response_code = 400

        return response_data, response_code

    def delete(self, user_id: int) -> tuple:
        response = {
            'error': 'User doesn\'t exist',
        }
        status_code = 404

        user = UserModel.get_or_none(UserModel.id == user_id)

        if isinstance(user, UserModel):
            if user.deleted_at is None:
                user.deleted_at = datetime.utcnow()
                user.save()

                user_dict = user.serialize()

                response = {
                    'data': user_dict,
                }
                status_code = 200
            else:
                response = {
                    'error': 'User already deleted',
                }
                status_code = 400

        return response, status_code


@api.resource('/search')
class UsersResource(UserResource):
    def post(self) -> tuple:
        data = request.get_json()

        user_fields = UserModel.get_fields(['id'])
        v = Validator(schema=search_model_schema(user_fields))
        v.allow_unknown = False

        if not v.validate(data):
            return {
                       'message': 'validation error',
                       'fields': v.errors,
                   }, 422

        page_number, items_per_page, order_by = self.get_request_query_fields(data)

        query = UserModel.select()
        records_total = query.count()

        query = self.create_query(query, data)

        query = (query.order_by(order_by)
                 .paginate(page_number, items_per_page))

        records_filtered = query.count()
        user_list = []

        for user in query:
            user_dict = user.serialize()
            user_list.append(user_dict)

        return {
                   'data': user_list,
                   'records_total': records_total,
                   'records_filtered': records_filtered,
               }, 200


@api.resource('/xlsx')
class ExportUsersExcelResource(UserResource):
    def post(self) -> tuple:
        def write_excel_rows(rows: list, workbook: Workbook, worksheet: Worksheet) -> int:
            excel_longest_word = ''

            for i, row in enumerate(rows, 1):
                format = None

                if i == 1:
                    format = workbook.add_format({
                        'bold': True,
                        'bg_color': '#cccccc'
                    })
                elif i % 2 == 0:
                    format = workbook.add_format({
                        'bg_color': '#f1f1f1'
                    })

                range_cells = "A%s:I10" % i

                row_longest_word = find_longest_word(row)

                if len(row_longest_word) > len(excel_longest_word):
                    excel_longest_word = row_longest_word

                worksheet.write_row(range_cells, row, format)

            return len(excel_longest_word)

        def adjust_each_column_width(rows: list, worksheet: Worksheet, excel_longest_word: int) -> None:
            if rows:
                for i, v in enumerate(rows[0]):
                    formatted_row = [str(item) for item in v]
                    max_column_width = max(formatted_row, key=len)
                    max_column_width_len = len(max_column_width)

                    # worksheet.set_column(i, i + 1, max_column_width_len + 2)
                    worksheet.set_column(i, i + 1, excel_longest_word + 1)

        storage_dir = current_app.config.get('STORAGE_DIRECTORY')
        file_prefix = datetime.utcnow().strftime('%Y%m%d_%H%M%S')
        excel_filename = '{}/{}_users.xlsx'.format(storage_dir, file_prefix)

        rows = []

        workbook = xlsxwriter.Workbook(excel_filename)
        worksheet = workbook.add_worksheet()
        worksheet.set_zoom(120)

        original_column_names = UserModel.get_fields(['id'], self.column_display_order)
        self.format_column_names(rows, original_column_names)

        data = request.get_json()

        page_number, items_per_page, order_by = self.get_request_query_fields()

        select_fields = [
            UserModel._meta.fields[column_name]
            for column_name in original_column_names
        ]

        query = UserModel.select(*select_fields)

        query = self.create_query(query, data)

        users_query = (query.order_by(order_by)
                       .paginate(page_number, items_per_page))

        self.format_user_data(users_query, rows)

        total_fields = len(UserModel.get_fields(['id'])) - 1
        columns = 'A1:%s10' % pos_to_char(total_fields).upper()
        worksheet.autofilter(columns)

        excel_longest_word = write_excel_rows(rows, workbook, worksheet)
        adjust_each_column_width(rows, worksheet, excel_longest_word)

        workbook.close()

        kwargs = {
            'filename_or_fp': excel_filename,
            'mimetype': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            'as_attachment': True,
            'attachment_filename': excel_filename,
        }

        return send_file(**kwargs)


@api.resource('/pdf')
class ExportUsersPdfResource(UserResource):
    def post(self) -> tuple:
        def write_docx_content(rows: list, document: Document) -> None:
            header_fields = rows[0]

            table = document.add_table(rows=len(rows), cols=len(header_fields))

            for i in range(len(rows)):
                row = table.rows[i]
                for j, table_cell in enumerate(rows[i]):
                    row.cells[j].text = str(table_cell)

        storage_dir = current_app.config.get('STORAGE_DIRECTORY')
        file_prefix = datetime.utcnow().strftime('%Y%m%d_%H%M%S')
        basename = '{}_users'.format(file_prefix)
        filename = '{}.docx'.format(basename)
        docx_filename = '{}/{}'.format(storage_dir, filename)

        page_number, items_per_page, order_by = self.get_request_query_fields()
        rows = []

        original_column_names = UserModel.get_fields(['id'], self.column_display_order)
        self.format_column_names(rows, original_column_names)

        data = request.get_json()

        select_fields = [
            UserModel._meta.fields[column_name]
            for column_name in original_column_names
        ]

        query = UserModel.select(*select_fields)

        query = self.create_query(query, data)

        users_query = (query.order_by(order_by)
                       .paginate(page_number, items_per_page))

        self.format_user_data(users_query, rows)

        document = docx.Document()

        write_docx_content(rows, document)
        document.save(docx_filename)

        convert_to(storage_dir, docx_filename)

        pdf_filename = '{}/{}.pdf'.format(storage_dir, basename)

        kwargs = {
            'filename_or_fp': pdf_filename,
            'mimetype': 'application/pdf',
            'as_attachment': True,
            'attachment_filename': pdf_filename,
        }

        return send_file(**kwargs)
